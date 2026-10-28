"""Looking ahead and back in the books, on real PostgreSQL: the cash-flow forecast (feature 58),
staff pay (60), the year and zakat (62), the workbook for tax filing (63), the credit passport
(64), the seasons ahead (55) and the coach's weekly note (65)."""

from __future__ import annotations

import io
from datetime import date, timedelta
from typing import Any

import pytest
import respx
from openpyxl import load_workbook
from sqlalchemy import select, update

from chatfront.audit.service import SYSTEM
from chatfront.books import forecast, ledger, plan_routes
from chatfront.books.models import Category, EntryKind, LedgerEntry
from chatfront.channels.models import Conversation
from chatfront.core.i18n import t
from chatfront.core.models import Staff, StaffRole, Tenant
from chatfront.insights import coach
from chatfront.jobs.queue import enqueue
from chatfront.services import Services
from chatfront.shield.models import Consent
from tests.integration.chat import Client, fixture, owner_auth, shop_id, use
from tests.integration.test_books_shield import graph_mock, owner_on_whatsapp, texts
from tests.integration.test_couriers import packed_cod_order

NUMBER = "923004441122"


async def tenant_of(services: Services) -> Tenant:
    async with services.db.system() as session:
        tenant = await session.get(Tenant, await shop_id(session))
    assert tenant is not None
    return tenant


async def post(
    services: Services, kind: EntryKind, amount: int, day: date, category: Category, **extra: Any
) -> None:
    tenant = await tenant_of(services)
    async with services.db.tenant(tenant.id) as session:
        found = await session.get(Tenant, tenant.id)
        assert found is not None
        await ledger.post(
            session,
            found,
            kind,
            amount,
            SYSTEM,
            category=category,
            source=extra.pop("source", "console"),
            occurred_on=day,
            **extra,
        )


def month_day(today: date, back: int, day: int) -> date:
    year, month = today.year, today.month - back
    while month <= 0:
        year, month = year - 1, month + 12
    return date(year, month, day)


# -- cash-flow forecast (58) ------------------------------------------------------------------


FIXED_TODAY = date(2026, 10, 28)  # a Wednesday: the weeks ahead start on 2, 9, 16, 23 November


def fixed_day(monkeypatch: pytest.MonkeyPatch) -> date:
    monkeypatch.setattr(forecast, "today", lambda _tenant: FIXED_TODAY)
    monkeypatch.setattr(plan_routes, "today", lambda _tenant: FIXED_TODAY)
    return FIXED_TODAY


async def test_the_forecast_adds_up_the_weeks_ahead(
    app: tuple[Client, Services], monkeypatch: pytest.MonkeyPatch
) -> None:
    client, services = app
    owner = await owner_auth(services)
    today = fixed_day(monkeypatch)
    this_week = today - timedelta(days=today.weekday())
    for week in range(1, 9):
        monday = this_week - timedelta(weeks=week)
        await post(services, EntryKind.SALE, 60_000, monday, Category.SALES)
        await post(services, EntryKind.EXPENSE, 20_000, monday + timedelta(days=2), Category.STOCK)
    for back in (1, 2, 3):
        await post(services, EntryKind.EXPENSE, 150_000, month_day(today, back, 5), Category.RENT)
    await post(services, EntryKind.EXPENSE, 9_000, month_day(today, 1, 10), Category.UTILITIES)

    found = (await client.get("/api/owner/books/forecast", headers=owner)).json()
    assert (found["usual_in"], found["usual_out"]) == (60_000, 20_000)
    assert {(f["category"], f["amount"], f["day"]) for f in found["fixed"]} == {
        ("rent", 150_000, 5)
    }  # bills paid in only one month are not fixed yet
    assert found["cash_now"] is None
    assert all(w["balance"] is None for w in found["weeks"])
    assert len(found["weeks"]) == forecast.WEEKS_AHEAD

    counted = await client.put("/api/owner/books/cash", json={"amount": 20_000}, headers=owner)
    weeks = counted.json()["weeks"]
    balance = counted.json()["cash_now"]
    assert balance == 20_000
    for week in weeks:
        rent = sum(f["amount"] for f in week["fixed"])
        assert week["money_out"] == 20_000 + rent
        balance += week["money_in"] - week["money_out"]
        assert (week["balance"], week["short"]) == (balance, balance < 0)
    rent_weeks = [w["starts"] for w in weeks if w["fixed"]]
    assert rent_weeks == ["2026-11-02"]  # 5 November

    # Money that came in after the count moves the cash on.
    await post(services, EntryKind.SALE, 5_000, today, Category.SALES)
    again = (await client.get("/api/owner/books/forecast", headers=owner)).json()
    assert again["cash_now"] == 25_000

    # Staff pay set in the console replaces salaries seen in the books.
    me = (await client.get("/api/owner/team", headers=owner)).json()[0]["id"]
    await client.put(
        f"/api/owner/team/{me}/pay", json={"salary": 80_000, "pay_day": 1}, headers=owner
    )
    paid = (await client.get("/api/owner/books/forecast", headers=owner)).json()
    assert ("salaries", 80_000, 1) in {
        (f["category"], f["amount"], f["day"]) for f in paid["fixed"]
    }


async def test_the_owner_hears_of_a_short_week_on_monday(
    app: tuple[Client, Services], monkeypatch: pytest.MonkeyPatch
) -> None:
    client, services = app
    owner = await owner_auth(services)
    tenant = await tenant_of(services)
    today = fixed_day(monkeypatch)
    for back in (1, 2):
        await post(services, EntryKind.EXPENSE, 500_000, month_day(today, back, 3), Category.RENT)
    await client.put("/api/owner/books/cash", json={"amount": 1_000}, headers=owner)
    with respx.mock:
        graph = graph_mock()
        await owner_on_whatsapp(client, services, NUMBER)
        before = len(graph.calls)
        async with services.db.system() as session:
            await enqueue(session, forecast.SHOP_JOB, {}, tenant_id=tenant.id)
        assert services.worker is not None
        while await services.worker.run_once():
            pass
        sent = [m["text"]["body"] for m in texts(graph, before) if m["type"] == "text"]
    found = (await client.get("/api/owner/books/forecast", headers=owner)).json()
    short = next(w for w in found["weeks"] if w["short"])
    expected = forecast.short_text(
        tenant,
        forecast.Week(
            date.fromisoformat(short["starts"]),
            short["money_in"],
            short["money_out"],
            tuple(forecast.Fixed(f["category"], f["amount"], 3) for f in short["fixed"]),
            short["balance"],
        ),
        "ur-Latn",
    )
    assert short["starts"] == "2026-11-02"
    assert expected in sent


# -- staff pay (60) ---------------------------------------------------------------------------


async def test_staff_are_paid_salary_and_commission_once(app: tuple[Client, Services]) -> None:
    client, services = app
    owner = await owner_auth(services)
    added = await client.post(
        "/api/owner/team", json={"name": "Ali", "role": "sales"}, headers=owner
    )
    ali = added.json()["member"]["id"]
    bad = await client.put(
        f"/api/owner/team/{ali}/pay", json={"salary": 1, "commission_bp": 9000}, headers=owner
    )
    assert bad.status_code == 422
    await client.put(
        f"/api/owner/team/{ali}/pay",
        json={"salary": 3_000_000, "commission_bp": 200, "pay_day": 5},
        headers=owner,
    )
    tenant = await tenant_of(services)
    month = ledger.today(tenant).replace(day=1)
    assert ali  # the lawn shop's own staff are set up the same way

    # A delivered order in the grocer's shop came through a chat handed to Sana.
    grocery_owner = await owner_auth(services, "barkat-kiryana")
    first = await packed_cod_order(client, services, grocery_owner, "0300 7771111")
    url = f"/api/owner/orders/{first.code}/status"
    for step in ("dispatched", "delivered"):
        await client.post(url, json={"to": step}, headers=grocery_owner)
    async with services.db.system() as session:
        grocer = await shop_id(session, "barkat-kiryana")
        sana = Staff(tenant_id=grocer, name="Sana", email="sana@x.test", role=StaffRole.SALES)
        session.add(sana)
        await session.flush()
        await session.execute(
            update(Conversation)
            .where(Conversation.id == first.conversation_id)
            .values(assigned_staff_id=sana.id)
        )
        sana_id = sana.id
    await client.put(
        f"/api/owner/team/{sana_id}/pay",
        json={"salary": 2_000_000, "commission_bp": 500, "pay_day": 1},
        headers=grocery_owner,
    )
    pay = (
        await client.get(f"/api/owner/team/pay?month={month:%Y-%m}", headers=grocery_owner)
    ).json()
    row = next(p for p in pay if p["name"] == "Sana")
    assert row["sales"] == first.total
    assert row["commission"] == first.total * 500 // 10_000
    assert row["due"] == 2_000_000 + row["commission"]
    paid = await client.post(f"/api/owner/team/{sana_id}/pay/{month:%Y-%m}", headers=grocery_owner)
    assert paid.json()["paid"] == row["due"]
    twice = await client.post(f"/api/owner/team/{sana_id}/pay/{month:%Y-%m}", headers=grocery_owner)
    assert twice.status_code == 409
    async with services.db.system() as session:
        entry = await session.scalar(
            select(LedgerEntry).where(
                LedgerEntry.tenant_id == grocer, LedgerEntry.category == Category.SALARIES
            )
        )
    assert entry is not None
    assert entry.amount == row["due"]
    nothing = next(p for p in pay if p["name"] != "Sana" and p["salary"] == 0)
    none_due = await client.post(
        f"/api/owner/team/{nothing['staff_id']}/pay/{month:%Y-%m}", headers=grocery_owner
    )
    assert none_due.json()["error"]["code"] == "nothing_due"
    bad_month = await client.get("/api/owner/team/pay?month=2026-13", headers=grocery_owner)
    assert bad_month.status_code == 400


# -- the year, zakat and the workbook (62, 63) -------------------------------------------------


async def test_the_year_zakat_and_the_tax_workbook(app: tuple[Client, Services]) -> None:
    client, services = app
    owner = await owner_auth(services)
    tenant = await tenant_of(services)
    today = ledger.today(tenant)
    await post(services, EntryKind.SALE, 485_000, today, Category.SALES)
    await post(services, EntryKind.EXPENSE, 100_000, today, Category.RENT, description="Shop rent")
    await post(services, EntryKind.REFUND, 20_000, today, Category.SALES)
    found = (await client.get("/api/owner/books/year", headers=owner)).json()
    start = date.fromisoformat(found["start"])
    assert (start.month, start.day) == (7, 1)  # Pakistan's tax year
    assert found["label"] == f"{start.year}-{(start.year + 1) % 100:02d}"
    assert (found["sales"], found["expenses"], found["refunds"]) == (485_000, 100_000, 20_000)
    assert found["cash_profit"] == 365_000
    assert found["expense_categories"] == {"rent": 100_000}
    assert found["months"][-1]["money_in"] == 485_000
    stock = found["zakat_inputs"]["stock"]
    assert stock > 0  # the demo products' opening batches have a cost

    zakat = await client.post(
        "/api/owner/books/zakat",
        json={"cash": 4_000_000, "debts": 1_000_000, "nisab": 1_000},
        headers=owner,
    )
    figures = zakat.json()
    assert figures["net"] == 4_000_000 + stock + figures["owed_to_shop"] - 1_000_000
    assert figures["due"] == figures["net"] // 40
    below = await client.post(
        "/api/owner/books/zakat",
        json={"cash": 0, "stock": 0, "owed_to_shop": 0, "nisab": 1_000_000_000},
        headers=owner,
    )
    assert below.json()["due"] == 0

    sheet = await client.get(
        f"/api/owner/books/export.xlsx?start={today - timedelta(days=30)}&end={today}",
        headers=owner,
    )
    assert sheet.headers["content-type"].startswith("application/vnd.openxmlformats")
    book = load_workbook(io.BytesIO(sheet.content))
    assert book.sheetnames == ["Summary", "Sales", "Expenses", "Refunds"]
    assert book["Summary"]["B4"].value == 4850
    assert [c.value for c in book["Expenses"][2]][1:4] == ["rent", 1000, None]
    assert book["Refunds"].max_row == 2
    too_long = await client.get(
        "/api/owner/books/export.xlsx?start=2020-01-01&end=2026-01-01", headers=owner
    )
    assert too_long.json()["error"]["code"] == "bad_range"


# -- the credit passport (64) -----------------------------------------------------------------


async def test_a_credit_passport_is_signed_shared_and_revoked(
    app: tuple[Client, Services],
) -> None:
    client, services = app
    owner = await owner_auth(services)
    wording = "I agree to share my shop's sales history with the lender I give this code to."
    refused = await client.post(
        "/api/owner/passports", json={"agree": False, "wording": wording}, headers=owner
    )
    assert refused.json()["error"]["code"] == "not_agreed"
    made = (
        await client.post(
            "/api/owner/passports", json={"agree": True, "wording": wording}, headers=owner
        )
    ).json()
    assert made["code"].startswith("CP")
    assert set(made["totals"]) >= {"orders_delivered", "paid_matched_to_sms", "cash_at_door"}
    async with services.db.system() as session:
        agreed = await session.scalar(select(Consent).where(Consent.purpose == "credit_passport"))
    assert agreed is not None
    assert agreed.wording == wording

    public = (await client.get(f"/api/passport/{made['code'].lower()}")).json()
    assert public["valid"] is True
    assert public["revoked"] is False
    assert len(public["statement"]["months"]) == 12
    assert public["statement"]["shop"] == (await tenant_of(services)).name

    revoked = await client.post(f"/api/owner/passports/{made['code']}/revoke", headers=owner)
    assert revoked.json() == {"revoked": True}
    assert (await client.get(f"/api/passport/{made['code']}")).json()["revoked"] is True
    assert (
        await client.post(f"/api/owner/passports/{made['code']}/revoke", headers=owner)
    ).status_code == 404
    assert (await client.get("/api/passport/NOSUCH")).status_code == 404
    listed = (await client.get("/api/owner/passports", headers=owner)).json()
    assert [p["revoked"] for p in listed] == [True]


# -- seasons and the coach (55, 65) -------------------------------------------------------------


async def test_the_seasons_ahead_and_the_coach_note(app: tuple[Client, Services]) -> None:
    client, services = app
    owner = await owner_auth(services)
    seasons = (await client.get("/api/owner/seasons", headers=owner)).json()
    assert all(s["starts"] <= s["ends"] for s in seasons)
    use(
        services,
        fixture(
            "business_coach",
            {"advice": ["Post your lawn suits again this week.", "Sell 3 for 2 on dupattas."]},
        ),
    )
    note = (await client.post("/api/owner/coach/now", headers=owner)).json()
    assert note["advice"] == ["Post your lawn suits again this week."]  # a number is dropped
    assert note["lines"]
    again = (await client.post("/api/owner/coach/now", headers=owner)).json()
    assert again["created_at"] == note["created_at"]  # one note a week
    listed = (await client.get("/api/owner/coach", headers=owner)).json()
    assert len(listed) == 1
    tenant = await tenant_of(services)
    async with services.db.tenant(tenant.id) as session:
        found = await coach.facts(session, tenant, ledger.today(tenant))
    assert len(coach.lines(found, "en")) == len(note["lines"])  # one line per fact, any language
    quiet = coach.Facts(week=found.week)
    assert coach.lines(quiet, "en") == [t("coach.steady", "en")]
