"""Restock from a sheet (feature 117; feature 85 reuses it for catalog imports): a CSV or Excel
file from the console. Spark Open says which column is which; without a model, code matches
the headers by their usual names. Then every row is matched to a product like a bill line,
and the owner checks the rows, fixes the ones that did not match, and confirms.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
import zipfile
from collections.abc import Sequence
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from pydantic import Field

from chatfront.ai.client import AiClient
from chatfront.ai.prompts import load_prompt
from chatfront.ai.types import ChatMessage, ChatRequest, Role
from chatfront.catalog.drafts import PurchaseLine, Strict
from chatfront.catalog.service import CatalogItem
from chatfront.core.errors import AppError
from chatfront.core.numbers import find_numbers
from chatfront.stock import restock

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
SAMPLE_ROWS = 3
FIELDS = ("item", "option", "quantity", "unit_cost", "total_cost", "expiry", "supplier")
# Usual header words per field, in the languages shops write sheets in.
HEADER_WORDS: dict[str, tuple[str, ...]] = {
    "item": ("item", "product", "name", "article", "description", "maal", "cheez", "نام", "آئٹم"),
    "option": ("colour", "color", "size", "variant", "rang", "رنگ", "سائز"),
    "quantity": ("qty", "quantity", "pcs", "pieces", "units", "tadaad", "تعداد", "مقدار"),
    "unit_cost": ("unit cost", "rate", "cost price", "unit price", "per piece", "ریٹ", "قیمت"),
    "total_cost": ("total", "amount", "line total", "raqam", "رقم", "کل"),
    "expiry": ("expiry", "exp", "expires", "best before", "میعاد"),
    "supplier": ("supplier", "vendor", "party", "سپلائر"),
}
_DMY = re.compile(r"^\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\s*$")


class Columns(Strict):
    """Which header holds what; null when no column does."""

    item: str | None = Field(description="The header of the product name column")
    option: str | None = None
    quantity: str | None = None
    unit_cost: str | None = None
    total_cost: str | None = None
    expiry: str | None = None
    supplier: str | None = None


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(data: bytes) -> list[list[str]]:
    """The sheet's rows as text, header first: an Excel workbook's first sheet, or a CSV."""
    if len(data) > MAX_BYTES:
        raise AppError("sheet_too_big", "A sheet can be up to 2 MB", status_code=413)
    if zipfile.is_zipfile(io.BytesIO(data)):
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise AppError("bad_sheet", "This Excel file could not be read") from exc
        sheet = book.worksheets[0]
        rows = [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect: Any = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2:
        raise AppError("empty_sheet", "The sheet needs a header row and at least one line")
    return rows[: MAX_ROWS + 1]


def keyword_columns(headers: Sequence[str]) -> Columns:
    """Code's own guess from the header words."""
    found: dict[str, str | None] = dict.fromkeys(FIELDS)
    for header in headers:
        said = header.casefold()
        for field in FIELDS:
            if found[field] is None and any(word in said for word in HEADER_WORDS[field]):
                found[field] = header
                break
    return Columns(**found)


def checked(columns: Columns, headers: Sequence[str]) -> Columns:
    """Only headers that really exist; an item and a quantity column are needed."""
    real = set(headers)
    kept = {f: v if v in real else None for f, v in columns.model_dump().items()}
    if not kept["item"] or not kept["quantity"]:
        raise AppError("no_columns", "The sheet needs a product column and a quantity column")
    return Columns(**kept)


async def map_columns(
    ai: AiClient, tenant_id: uuid.UUID, rows: Sequence[Sequence[str]], lang: str
) -> Columns:
    headers = list(rows[0])
    try:
        prompt = load_prompt("map_columns")
        sample = "\n".join(" | ".join(row) for row in rows[: SAMPLE_ROWS + 1])
        request = ChatRequest(
            role=Role.CHAT,
            messages=[ChatMessage("system", prompt.render()), ChatMessage("user", sample)],
            tenant_id=tenant_id,
            prompt=prompt.ref,
            lang=lang,
        )
        columns, _ = await ai.complete_json(request, Columns)
        return checked(columns, headers)
    except AppError:  # no model, or an answer that does not fit: code reads the headers
        return checked(keyword_columns(headers), headers)


def _iso(text: str) -> str | None:
    if restock.parse_expiry(text):
        return text.strip()
    found = _DMY.match(text)
    if not found:
        return None
    try:
        return date(int(found[3]), int(found[2]), int(found[1])).isoformat()
    except ValueError:
        return None


def _number(text: str) -> float | None:
    numbers = find_numbers(text)
    return float(numbers[0]) if numbers else None


def to_lines(
    rows: Sequence[Sequence[str]], columns: Columns
) -> tuple[list[PurchaseLine], str | None]:
    """Purchase lines from the rows, and the supplier if the sheet names one."""
    index = {header: i for i, header in enumerate(rows[0])}

    def get(row: Sequence[str], field: str) -> str:
        header = getattr(columns, field)
        position = index.get(header) if header else None
        return row[position] if position is not None and position < len(row) else ""

    lines, supplier = [], None
    for row in rows[1:]:
        item = " ".join(part for part in (get(row, "item"), get(row, "option")) if part)
        if not item:
            continue
        quantity = _number(get(row, "quantity"))
        lines.append(
            PurchaseLine(
                item=item[:160],
                quantity=int(quantity) if quantity and quantity.is_integer() else None,
                unit_cost=_number(get(row, "unit_cost")),
                total_cost=_number(get(row, "total_cost")),
                expiry=_iso(get(row, "expiry")),
            )
        )
        supplier = supplier or (get(row, "supplier") or None)
    return lines, supplier


def preview(
    lines: Sequence[PurchaseLine], items: Sequence[CatalogItem], digits: int
) -> list[dict[str, Any]]:
    """Each row as the console shows it: matched to a variant, or what is wrong with it."""
    rows = []
    for number, line in enumerate(lines, start=1):
        ops, problems = restock.resolve_lines([line], items, digits)
        row: dict[str, Any] = {
            "row": number,
            "item": line.item,
            "quantity": line.quantity,
            "total_cost": restock.line_cost(line, digits),
            "expires_on": line.expiry,
            "variant_id": None,
            "name": None,
            "problem": None,
        }
        if ops:
            row |= {"variant_id": ops[0]["variant_id"], "name": ops[0]["name"]}
        elif problems:
            row["problem"] = problems[0][0]
        rows.append(row)
    return rows
